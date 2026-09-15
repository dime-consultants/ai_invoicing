# tools/reconciliation_report.py
"""
Shared formatting for the "outstanding items" reconciliation report — the
client's standard bank-reconciliation summary layout: four sections
(Outstanding Ledger credits/debits, Outstanding Statement credits/debits),
each with Date/Ref/Item Description/Amount columns and a Total row.

Used by two callers that must stay byte-for-byte consistent:
  - tools/views.py::tool_call_report_download — builds this on the fly from
    ANY successful ToolCall whose result looks like a reconcile_datasets
    output (ad hoc "download report" from the Uploads/Tools UI).
  - tools/handlers.py::write_reconciliation_report — a builtin tool the chat
    agent calls directly so the file it attaches to the conversation is this
    same deterministic layout, not a free-form table it composes itself.
"""

from __future__ import annotations

RECONCILIATION_STATUSES = {"MATCH", "VARIANCE", "MISSING_IN_A", "MISSING_IN_B"}
RECONCILIATION_ACCT_FORMAT = '_(* #,##0.00_);_(* \\(#,##0.00\\);_(* "-"??_);_(@_)'

# (bucket key, section header, italic subtitle) — mirrors a hand-built bank
# reconciliation "Summary" sheet: outstanding items are grouped by which
# side they're missing from (Ledger = A, Statement = B), then split again
# by credit (positive amount) vs debit (negative amount).
RECONCILIATION_LAYOUT = [
    ("ledger_credits",    "Outstanding Ledger credits",    "Credits in A that are missing a debit in B"),
    ("ledger_debits",     "Outstanding Ledger Debits",     "Debits in A that are missing a credit in B"),
    ("statement_credits", "Outstanding Statement credits", "Credits in B that are missing a debit in A"),
    ("statement_debits",  "Outstanding Statement Debits",  "Debits in B that are missing a credit in A"),
]


def is_reconciliation_result(data) -> bool:
    rows = data.get("rows") if isinstance(data, dict) else None
    return isinstance(rows, list) and any(
        isinstance(r, dict) and r.get("status") in RECONCILIATION_STATUSES for r in rows
    )


def reconciliation_sections(data: dict) -> dict:
    """Split reconcile_datasets' flat MISSING_IN_A/MISSING_IN_B rows into
    the four outstanding buckets, by which side they're on and the sign of
    their amount. MATCH/VARIANCE rows aren't "outstanding" on either side,
    so they're intentionally excluded here — this is specifically the
    unreconciled-items summary, not a full dump of every row.
    """
    sections: dict = {key: [] for key, _, _ in RECONCILIATION_LAYOUT}
    for row in data.get("rows") or []:
        if not isinstance(row, dict):
            continue
        status_val = row.get("status")
        if status_val == "MISSING_IN_B":
            amt, bucket_pair = row.get("amount_a"), ("ledger_credits", "ledger_debits")
        elif status_val == "MISSING_IN_A":
            amt, bucket_pair = row.get("amount_b"), ("statement_credits", "statement_debits")
        else:
            continue
        if amt is None:
            continue
        bucket = bucket_pair[0] if amt >= 0 else bucket_pair[1]
        sections[bucket].append((row.get("id"), abs(amt)))
    return sections


def reconciliation_title(label_a: str, label_b: str) -> str:
    return f"Taking A is {label_a or 'File A'} and B is {label_b or 'File B'}"


def build_reconciliation_xlsx(title: str, sections: dict) -> bytes:
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    italic_subtitle = Font(italic=True, size=9)

    r = 1
    ws.cell(row=r, column=1, value=title)
    r += 2

    for key, header, subtitle in RECONCILIATION_LAYOUT:
        ws.cell(row=r, column=1, value=header)
        r += 1
        ws.cell(row=r, column=1, value=subtitle).font = italic_subtitle
        r += 1
        for col, label in enumerate(["Date", "Ref", "Item Description", "Amount"], start=1):
            ws.cell(row=r, column=col, value=label)
        r += 1
        first_data_row = r
        for ref, amount in sections[key]:
            ws.cell(row=r, column=2, value=ref)
            ws.cell(row=r, column=4, value=amount).number_format = RECONCILIATION_ACCT_FORMAT
            r += 1
        last_data_row = r - 1
        ws.cell(row=r, column=1, value="Total")
        total_cell = ws.cell(row=r, column=4)
        total_cell.value = (
            f"=SUM(D{first_data_row}:D{last_data_row})" if last_data_row >= first_data_row else 0
        )
        total_cell.number_format = RECONCILIATION_ACCT_FORMAT
        r += 2

    ws.column_dimensions["C"].width = 16

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def reconciliation_text_rows(title: str, sections: dict) -> list:
    """Same section layout as build_reconciliation_xlsx, flattened to plain
    row-arrays for csv/pdf renderers that can't do per-cell formulas/styling.
    """
    out = [[title], []]
    for key, header, subtitle in RECONCILIATION_LAYOUT:
        out.append([header])
        out.append([subtitle])
        out.append(["Date", "Ref", "Item Description", "Amount"])
        total = 0
        for ref, amount in sections[key]:
            out.append(["", ref, "", f"{amount:.2f}"])
            total += amount
        out.append(["Total", "", "", f"{total:.2f}"])
        out.append([])
    return out


def build_reconciliation_csv(title: str, sections: dict) -> bytes:
    import csv as csv_module
    import io

    buf = io.StringIO()
    writer = csv_module.writer(buf)
    for row in reconciliation_text_rows(title, sections):
        writer.writerow(row)
    return buf.getvalue().encode("utf-8")


def build_reconciliation_pdf(title: str, sections: dict) -> bytes:
    import io

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter, topMargin=0.5 * inch, bottomMargin=0.5 * inch)
    elements = [Paragraph(title, styles["Title"]), Spacer(1, 0.2 * inch)]

    for key, header, subtitle in RECONCILIATION_LAYOUT:
        elements.append(Paragraph(header, styles["Heading3"]))
        elements.append(Paragraph(subtitle, styles["Italic"]))
        table_data = [["Date", "Ref", "Item Description", "Amount"]]
        total = 0
        for ref, amount in sections[key]:
            table_data.append(["", str(ref) if ref is not None else "", "", f"{amount:,.2f}"])
            total += amount
        table_data.append(["Total", "", "", f"{total:,.2f}"])
        table = Table(table_data, repeatRows=1)
        table.setStyle(TableStyle([
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EEF2F7")),
            ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ]))
        elements.append(table)
        elements.append(Spacer(1, 0.25 * inch))

    doc.build(elements)
    return buf.getvalue()
